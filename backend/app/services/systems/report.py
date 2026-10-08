"""SB2-107: the reviews and findings report (CONTRACTS_P2 §8.6), as a workbook or one CSV.

Pure rendering of a redacted live document and its redacted open reviews: what a reader sees in
the Changes and Findings trays, one row per review item and per finding, so the counts match.
"""

from __future__ import annotations

import csv
import io
from typing import Any, Mapping, Optional, Sequence

from app.services.systems.drift import pad_sort_key
from app.services.systems.icd import _end_label

RENDERER_VERSION = "1"

ITEM_LABELS = {
    "connector_missing": "Connector missing",
    "connector_changed": "Connector changed",
    "pin_missing": "Pin missing",
    "net_changed": "Net changed",
    "signal_mismatch": "Signal does not match",
}
REVIEW_LABELS = {
    "source_update": "Board change",
    "baseline_unreachable": "Baseline unreachable",
    "child_update": "Subsystem or module update",
    "import": "CSV import",
    "manifest_import": "Manifest changed outside Prism",
}
_SEVERITY_ORDER = {"error": 0, "warning": 1, "info": 2}

REVIEW_COLUMNS = ("Review", "Kind", "Board", "From", "To", "Opened", "Item", "Change", "Link", "End",
                  "Connector", "Pins", "Expected", "Observed", "Decision")
FINDING_COLUMNS = ("Severity", "Rule", "Finding", "Board", "Connector", "Pin", "Link", "State", "Note",
                   "Waived by", "Waived on")


def _nets(value: Any) -> str:
    """A pin's net set as text: ``(no net)`` for an unconnected pin, nothing when unknown."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return " | ".join(str(v) for v in value) if value else "(no net)"
    if isinstance(value, Mapping):
        if "libId" in value or "footprint" in value:
            parts = [value.get("reference"), value.get("libId"), value.get("footprint")]
            return " · ".join(str(p) for p in parts if p)
        if "leaves" in value:
            return _nets(value["leaves"])
        return ", ".join(f"{k}: {_nets(v) if isinstance(v, (list, tuple, Mapping)) else v}" for k, v in value.items())
    return str(value)


def _proposal(value: Any) -> str:
    """An import item's proposed row or wire, as the Changes tray words it."""
    if not isinstance(value, Mapping):
        return _nets(value)

    def side(name: str) -> str:
        entry = value.get(name) or {}
        pad = f"{entry.get('label', '')}/{entry.get('reference', '')}.{entry.get('pin', '')}" if entry else ""
        if value.get("kind") != "wire":
            return pad
        ordinal, pin = value.get(f"{name}End"), value.get(f"{name}Pin")
        end = f"End {'?' if ordinal is None else ordinal + 1}.{pin or '?'}"
        return f"{end} ({pad})" if pad else f"{end} (not mated)"

    text = f"{side('from')} ↔ {side('to')}"
    if value.get("signal"):
        text += f" · signal {value['signal']}"
    return f"{value['linkName']}: {text}" if value.get("kind") == "wire" and value.get("linkName") else text


def _date(iso: Optional[str]) -> str:
    return (iso or "")[:10]


def _short(commit: Optional[str]) -> str:
    return (commit or "")[:12]


def _names(document: Mapping[str, Any]) -> tuple[dict[str, str], dict[str, Mapping[str, Any]], dict[str, str]]:
    labels = {i["id"]: i["label"] for i in document["instances"]}
    links = {link["id"]: link for link in document["links"]}
    names = {link["id"]: link["name"] or " ↔ ".join(_end_label(labels, link[e]) for e in ("a", "b"))
             for link in document["links"]}
    names.update({h["id"]: h["name"] for h in document.get("harnesses") or []})
    return labels, links, names


def review_rows(document: Mapping[str, Any], reviews: Sequence[Mapping[str, Any]]) -> list[list[str]]:
    """One row per open review item; a review the reader cannot see is one ``Restricted`` row."""
    labels, links, names = _names(document)
    out: list[list[str]] = []
    for review in sorted(reviews, key=lambda r: (r["createdAt"] or "", r["id"])):
        board = labels.get(review["instanceId"] or "", "")
        head = [review["id"], REVIEW_LABELS.get(review["kind"], review["kind"]), board,
                _short(review.get("fromCommit")), _short(review.get("toCommit")), _date(review["createdAt"])]
        if review.get("redacted"):
            out.append(head + ["", "Restricted"] + [""] * 7)
            continue
        items = review.get("items") or []
        if not items:
            summary = ((review.get("pendingChanges") or {}).get("summary")) if review["kind"] == "manifest_import" else None
            areas = ", ".join(sorted(summary)) if isinstance(summary, Mapping) else ""
            out.append(head + ["", "Whole review", "", "", "", "", "", areas, ""])
            continue
        for item in sorted(items, key=lambda i: i["ordinal"]):
            if item.get("redacted"):
                out.append(head + [str(item["ordinal"]), "Restricted"] + [""] * 6
                           + [item["decision"] or ""])
                continue
            link = links.get(item.get("linkId") or "")
            end = item.get("end")
            connector = _end_label(labels, link[end]) if link and end else ""
            pins = ", ".join(sorted(item.get("pins") or [], key=pad_sort_key))
            if review["kind"] == "import":
                expected, observed = _nets(item.get("expected")), _proposal(item.get("observed"))
            else:
                expected, observed = _nets(item.get("expected")), _nets(item.get("observed"))
            decision = item["decision"] or ""
            pad = (item.get("decisionPayload") or {}).get("pad")
            if decision and pad:
                decision += f" → {pad}"
            out.append(head + [str(item["ordinal"]), ITEM_LABELS.get(item["kind"], item["kind"]),
                               names.get(item.get("linkId") or "", ""), (end or "").upper(), connector, pins,
                               expected, observed, decision])
    return out


def finding_rows(document: Mapping[str, Any]) -> list[list[str]]:
    """Every finding, open ones first by severity, then the waived; then waivers no longer raised."""
    labels, _links, names = _names(document)
    validation = document.get("validation") or {}
    findings = validation.get("findings") or []

    def row(finding: Mapping[str, Any]) -> list[str]:
        waived = finding.get("waived")
        link_id = finding.get("linkId") or (finding.get("detail") or {}).get("harnessId") or ""
        return [finding["severity"], finding["rule"], finding["name"].replace("_", " "),
                labels.get(finding.get("instanceId") or "", ""), finding.get("reference") or "",
                finding.get("pin") or "", names.get(link_id, ""), "waived" if waived else "open",
                (waived or {}).get("note") or "", ((waived or {}).get("by") or "").removeprefix("user:"),
                _date((waived or {}).get("at"))]

    def order(finding: Mapping[str, Any]) -> tuple:
        return (bool(finding.get("waived")), _SEVERITY_ORDER.get(finding["severity"], 3), finding["rule"],
                labels.get(finding.get("instanceId") or "", ""), finding.get("reference") or "",
                pad_sort_key(finding.get("pin") or ""))

    out = [row(f) for f in sorted(findings, key=order)]
    for waiver in validation.get("waivers") or []:
        if waiver.get("active"):
            continue
        parts = ((waiver.get("findingKey") or "").split("|") + [""] * 7)[:7]
        out.append(["", waiver.get("rule") or "", "", labels.get(parts[1], ""), parts[5], parts[6], "",
                    "no longer raised", waiver.get("note") or "", (waiver.get("by") or "").removeprefix("user:"),
                    _date(waiver.get("at"))])
    return out


def summary_rows(document: Mapping[str, Any], reviews: Sequence[Mapping[str, Any]], *,
                 version: int, generated_at: str) -> list[list[str]]:
    """The heading facts and the counts the workspace shows (top bar, Findings and Changes badges)."""
    labels = {i["id"]: i["label"] for i in document["instances"]}
    validation = document.get("validation") or {}
    counts = validation.get("counts") or {}
    items = [item for review in reviews for item in review.get("items") or []]
    rows = [
        ["System", document["system"]["name"]],
        ["Version", str(version)],
        ["Generated", generated_at],
        ["Errors", str(counts.get("error", 0))],
        ["Warnings", str(counts.get("warning", 0))],
        ["Info", str(counts.get("info", 0))],
        ["Waived", str(counts.get("waived", 0))],
        ["Open reviews", str(len(reviews))],
        ["Review items", str(len(items))],
        ["Undecided items", str(sum(1 for item in items if not item.get("decision")))],
    ]
    for entry in validation.get("notEvaluated") or []:
        rows.append(["Not evaluated", f"{entry['rule']} · {labels.get(entry.get('instanceId') or '', '')} · {entry['reason']}"])
    return rows


def sections(document: Mapping[str, Any], reviews: Sequence[Mapping[str, Any]], *, version: int,
             generated_at: str) -> list[tuple[str, Sequence[str], list[list[str]]]]:
    return [
        ("Summary", ("Item", "Value"), summary_rows(document, reviews, version=version, generated_at=generated_at)),
        ("Reviews", REVIEW_COLUMNS, review_rows(document, reviews)),
        ("Findings", FINDING_COLUMNS, finding_rows(document)),
    ]


def _cell(value: str) -> str:
    """CSV only: neutralise a cell a spreadsheet would read as a formula (CSV injection). The
    workbook stores text cells as text instead, so ``+3V3`` stays as typed there."""
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def render_csv(parts: Sequence[tuple[str, Sequence[str], list[list[str]]]]) -> str:
    """Each section as a titled block: its name, its header, its rows, then a blank line."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    for index, (title, columns, rows) in enumerate(parts):
        if index:
            writer.writerow([])
        writer.writerow([f"# {title}"])
        writer.writerow(columns)
        writer.writerows([[_cell(v) for v in row] for row in rows])
    return buffer.getvalue()


def render_xlsx(parts: Sequence[tuple[str, Sequence[str], list[list[str]]]]) -> bytes:
    """One sheet per section, header row bold and frozen, columns sized to their content."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)
    for title, columns, rows in parts:
        sheet = workbook.create_sheet(title)
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.font = bold
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"  # text as typed: "=…" or "+3V3" never becomes a formula
        sheet.freeze_panes = "A2"
        if rows and title != "Summary":
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        for index, column in enumerate(columns, start=1):
            width = max([len(column)] + [len(str(row[index - 1])) for row in rows[:2000] if index <= len(row)])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
