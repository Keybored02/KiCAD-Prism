"""SB2-107 (P2 §8.6): the reviews and findings report, as a workbook and as one CSV."""

from __future__ import annotations

import csv
import io
import unittest

from fastapi import FastAPI
from openpyxl import load_workbook
from unittest import mock

from test_system_api import _request
from test_system_snapshots import DESIGNER, VIEWER, SnapshotCase

from app.api import systems as systems_api
from app.services.systems import service as service_module


def sheets(content: bytes) -> dict[str, list[list[str]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    return {sheet.title: [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets}


def records(rows: list[list[str]]) -> list[dict[str, str]]:
    header, *body = rows
    return [dict(zip(header, row)) for row in body]


class ReportCase(SnapshotCase):
    def open_f1(self) -> None:
        """F1: J7.17 loses its net, so OBC-A's review has one net_changed item."""
        self.move_track("mini_obc", "F1")
        self.detector.check_instance(self.instances["OBC-A"])

    def fan_out(self) -> None:
        """A second link reusing a pin of the first raises SYS-V02 warnings."""
        link = next(link for link in self.service.document(DESIGNER, self.sid).body["links"] if link["rows"])
        ends = {side: {"instanceId": link[side]["instanceId"], "portKey": link[side]["port"]["portKey"]} for side in "ab"}
        second = self.service.create_link(DESIGNER, self.sid, self.version(), a=ends["a"], b=ends["b"],
                                          name="Fan-out", harness=None).body
        row = link["rows"][0]
        self.service.replace_rows(DESIGNER, self.sid, self.version(), second["id"],
                                  [{"pinA": row["pinA"], "pinB": row["pinB"], "signal": "=HYPERLINK(\"x\")"}])

    def workbook(self, caller=DESIGNER) -> dict[str, list[list[str]]]:
        content, _name, _version = self.service.report(caller, self.sid, "xlsx")
        return sheets(content)


class ReportTest(ReportCase):
    def test_one_sheet_per_section_and_the_counts_match_the_trays(self) -> None:
        self.open_f1()
        self.fan_out()
        book = self.workbook()
        self.assertEqual(list(book), ["Summary", "Reviews", "Findings", "Renames"])
        summary = dict(book["Summary"][1:])
        validation = self.service.validation_report(DESIGNER, self.sid).body
        counts = self.service.document(DESIGNER, self.sid).body["findingCounts"]
        self.assertEqual((summary["Errors"], summary["Warnings"], summary["Info"]),
                         (str(counts["error"]), str(counts["warning"]), str(counts["info"])))
        self.assertEqual(summary["Version"], str(self.version()))
        findings = records(book["Findings"])
        self.assertEqual(len(findings), len(validation["findings"]))
        self.assertEqual(sum(1 for f in findings if f["Severity"] == "warning" and f["State"] == "open"),
                         counts["warning"])
        reviews = self.service.list_reviews(DESIGNER, self.sid, "open")
        self.assertEqual(summary["Open reviews"], str(len(reviews)))
        self.assertEqual(summary["Review items"], str(sum(len(r["items"]) for r in reviews)))

    def test_a_review_item_lists_its_pin_with_the_expected_and_observed_nets(self) -> None:
        self.open_f1()
        [item] = records(self.workbook()["Reviews"])
        self.assertEqual((item["Kind"], item["Board"], item["Change"], item["Link"], item["Pins"]),
                         ("Board change", "OBC-A", "Net changed", "L-J7J4", "17"))
        self.assertEqual(item["Expected"], "PAYLOAD_RESET#")
        self.assertEqual(item["Observed"], "(no net)")
        self.assertEqual(item["Decision"], "")
        self.assertTrue(item["Connector"].startswith("OBC-A J7"))

    def test_a_waived_finding_carries_its_note(self) -> None:
        self.fan_out()
        warning = next(f for f in self.service.validation_report(DESIGNER, self.sid).body["findings"]
                       if f["severity"] == "warning")
        self.service.waive_finding(DESIGNER, self.sid, self.version(), warning["key"], "Shared on purpose")
        waived = [f for f in records(self.workbook()["Findings"]) if f["State"] == "waived"]
        self.assertEqual([(f["Rule"], f["Note"], f["Waived by"]) for f in waived],
                         [(warning["rule"], "Shared on purpose", DESIGNER.actor.removeprefix("user:"))])
        self.assertEqual(dict(self.workbook()["Summary"][1:])["Waived"], "1")

    def test_typed_text_is_never_a_formula(self) -> None:
        self.fan_out()
        content, _name, _version = self.service.report(DESIGNER, self.sid, "xlsx")
        workbook = load_workbook(io.BytesIO(content))
        cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
                 if isinstance(cell.value, str) and cell.value.startswith("=")]
        self.assertTrue(all(cell.data_type == "s" for cell in cells))
        text, _name, _version = self.service.report(DESIGNER, self.sid, "csv")
        self.assertNotIn(",=", text)

    def test_the_csv_holds_every_section_as_a_titled_block(self) -> None:
        self.open_f1()
        text, name, version = self.service.report(DESIGNER, self.sid, "csv")
        self.assertEqual(version, self.version())
        rows = list(csv.reader(io.StringIO(text)))
        titles = [row[0] for row in rows if row and row[0].startswith("# ")]
        self.assertEqual(titles, ["# Summary", "# Reviews", "# Findings", "# Renames"])
        start = rows.index(["# Reviews"])
        self.assertEqual(rows[start + 1][:3], ["Review", "Kind", "Board"])
        self.assertEqual(rows[start + 2][7], "Net changed")

    def test_a_review_on_a_board_the_reader_cannot_see_is_one_restricted_row(self) -> None:
        self.open_f1()
        self.conn.execute("INSERT INTO ws_folders (id, visibility_mode, allowed_roles)"
                          " VALUES ('fld_admins', 'roles', '[\"admin\"]')")
        self.conn.execute("UPDATE ws_projects SET folder_id = 'fld_admins' WHERE id = 'prj_obc'")
        self.conn.commit()
        book = self.workbook(VIEWER)
        [row] = records(book["Reviews"])
        self.assertEqual((row["Board"], row["Change"], row["From"], row["To"], row["Expected"]),
                         ("OBC-A", "Restricted", "", "", ""))
        self.assertNotIn("PAYLOAD_RESET#", repr(book))


class ReportApiTest(ReportCase):
    def setUp(self) -> None:
        super().setUp()
        patcher = mock.patch.object(service_module, "service", self.service)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.app = FastAPI()
        self.app.include_router(systems_api.router, prefix="/api/systems")

    def test_a_viewer_downloads_both_formats(self) -> None:
        self.open_f1()
        for fmt, media in (("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                           ("csv", "text/csv; charset=utf-8")):
            response = _request(self.app, "GET", f"/api/systems/{self.sid}/report.{fmt}", user="viewer")
            self.assertEqual(response.status, 200)
            self.assertEqual(response.headers["content-type"], media)
            self.assertIn(f'-report.{fmt}"', response.headers["content-disposition"])
            self.assertIn("etag", response.headers)
        self.assertEqual(_request(self.app, "GET", f"/api/systems/{self.sid}/report.pdf").status, 422)


if __name__ == "__main__":
    unittest.main()
