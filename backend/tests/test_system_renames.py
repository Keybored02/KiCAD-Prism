"""SB2-106 (CONTRACTS_P2 §23, D-P2-57): net rename proposals, closed by the board's own commit."""

from __future__ import annotations

import csv
import io
import unittest

from fastapi import FastAPI
from unittest import mock

from test_system_api import _request
from test_system_snapshots import DESIGNER, VIEWER, SnapshotCase

from app.api import systems as systems_api
from app.services.systems import renames, service as service_module
from app.services.systems.jobs import extract_and_store
from app.services.systems.store import Conflict, Invalid
from app.services.systems.visibility import etag


class RenameCase(SnapshotCase):
    def setUp(self) -> None:
        super().setUp()
        for step in ("F11.1", "F11.2"):
            extract_and_store(self.projects["prj_obc"], self.commits["mini_obc"][step], self.connect)

    def spi_row(self) -> dict:
        """OBC-A J7.3 on L-J7J4: the net fixture step F11.2 renames to SPI_CLK."""
        link = self.store.get_link(self.sid, self.links["L-J7J4"])
        return next(row for row in link["rows"] if row["pin_a"] == "3")

    def propose(self, net: str | None = None, name: str = "SPI_CLK", note: str = "Match the payload") -> dict:
        net = net or self.spi_row()["net_a"][0]
        return self.service.propose_rename(DESIGNER, self.sid, self.version(), instance_id=self.instances["OBC-A"],
                                           net=net, name=name, note=note).body

    def accept_all(self, review_id: str) -> None:
        review = next(r for r in self.service.list_reviews(DESIGNER, self.sid, "open") if r["id"] == review_id)
        for item in review["items"]:
            self.service.decide(DESIGNER, self.sid, self.version(), review_id, item["id"], "accept", None)

    def document(self) -> dict:
        return self.service.document(DESIGNER, self.sid, include_validation=True).body


class ProposeTest(RenameCase):
    def test_a_proposal_covers_every_row_on_the_net(self) -> None:
        net = self.spi_row()["net_a"][0]
        body = self.propose()
        self.assertEqual((body["net"], body["name"], body["state"], body["rows"]), (net, "SPI_CLK", "open", 1))
        [listed] = self.document()["renames"]
        self.assertEqual((listed["id"], listed["rows"]), (body["id"], 1))
        [event] = self.events("rename_proposed")
        self.assertEqual(event["payload"]["name"], "SPI_CLK")

    def test_bad_proposals_are_refused(self) -> None:
        net = self.spi_row()["net_a"][0]
        for name, wanted in (("A/B", "invalid_name"), ("has space", "invalid_name"),
                             (renames.leaf(net), "already named")):
            with self.subTest(name=name), self.assertRaisesRegex(Invalid, wanted):
                self.propose(name=name)
        with self.assertRaisesRegex(Invalid, "no row"):
            self.propose(net="/Nowhere/NET")
        self.propose()
        with self.assertRaisesRegex(Conflict, "rename_open"):
            self.propose(name="SPI_CK")

    def test_a_withdrawn_proposal_leaves_the_document(self) -> None:
        body = self.propose()
        self.service.withdraw_rename(DESIGNER, self.sid, self.version(), body["id"])
        self.assertEqual(self.document()["renames"], [])
        with self.assertRaisesRegex(Conflict, "rename_closed"):
            self.service.withdraw_rename(DESIGNER, self.sid, self.version(), body["id"])

    def test_a_v09_finding_names_the_proposal_for_its_net(self) -> None:
        self.move_track("mini_obc", "F8")  # J7.18's schematic net changes; accepted, it no longer matches PAY
        self.accept_all(self.detector.check_instance(self.instances["OBC-A"]).review_id)
        self.assertEqual(self.store.get_system(self.sid).get("optional_rules") or [], [])
        findings = [f for f in self.document()["validation"]["findings"] if f["rule"] == "SYS-V09"]
        self.assertTrue(findings, "SYS-V09 runs without opting in (D-P2-57)")
        link = self.store.get_link(self.sid, findings[0]["linkId"])
        row = next(r for r in link["rows"] if r["id"] == findings[0]["rowId"])
        instance_id = link["a_instance_id"]
        if self.store.get_instance(self.sid, instance_id).get("kind", "board") != "board" or not row["net_a"]:
            self.skipTest("the first V09 row has no board net on end A")
        body = self.service.propose_rename(DESIGNER, self.sid, self.version(), instance_id=instance_id,
                                           net=row["net_a"][0], name="RENAMED_NET").body
        [annotated] = [f for f in self.document()["validation"]["findings"]
                       if f["rule"] == "SYS-V09" and f["rowId"] == row["id"]]
        self.assertEqual(annotated["detail"]["rename"]["id"], body["id"])


class ApplyTest(RenameCase):
    def test_the_boards_commit_with_the_rename_closes_it_without_a_review(self) -> None:
        body = self.propose()
        self.move_track("mini_obc", "F11.1")
        first = self.detector.check_instance(self.instances["OBC-A"])
        self.assertEqual(first.outcome, "review_opened")  # F1: J7.17 lost its net, a real change
        self.accept_all(first.review_id)
        tip = self.move_track("mini_obc", "F11.2")
        second = self.detector.check_instance(self.instances["OBC-A"])
        self.assertEqual(second.outcome, "auto_advanced")
        self.assertEqual(self.reviews("OBC-A", "open"), [])
        [stored] = self.store.list_renames(self.sid, ("applied",))
        self.assertEqual((stored["id"], stored["closed_commit"]), (body["id"], tip))
        self.assertEqual(renames.leaf(self.spi_row()["net_a"][0]), "SPI_CLK")
        self.assertEqual(self.document()["renames"], [])
        [event] = self.events("rename_applied")
        self.assertEqual(event["payload"]["commit"], tip)

    def test_a_rename_that_comes_with_other_changes_closes_when_its_review_applies(self) -> None:
        self.propose()
        self.move_track("mini_obc", "F11.2")
        result = self.detector.check_instance(self.instances["OBC-A"])
        self.assertEqual(result.outcome, "review_opened")
        self.assertEqual(len(self.store.list_renames(self.sid)), 1)
        self.accept_all(result.review_id)
        self.assertEqual(len(self.store.list_renames(self.sid, ("applied",))), 1)

    def test_an_unrelated_net_change_is_still_reviewed(self) -> None:
        self.propose(name="SOMETHING_ELSE")
        self.move_track("mini_obc", "F11.1")
        self.accept_all(self.detector.check_instance(self.instances["OBC-A"]).review_id)
        self.move_track("mini_obc", "F11.2")
        self.assertEqual(self.detector.check_instance(self.instances["OBC-A"]).outcome, "review_opened")
        self.assertEqual(len(self.store.list_renames(self.sid)), 1)


class BoardPageTest(RenameCase):
    def test_the_board_page_lists_the_systems_and_their_proposals(self) -> None:
        self.propose()
        body = self.service.systems_using_project(VIEWER, "prj_obc")
        [system] = body["systems"]
        self.assertEqual(sorted(i["label"] for i in system["instances"]), ["OBC-A", "OBC-B"])
        [proposal] = system["renames"]
        self.assertEqual((proposal["board"], proposal["name"], proposal["connectors"]), ("OBC-A", "SPI_CLK", ["J7.3"]))
        rows = list(csv.DictReader(io.StringIO(self.service.project_renames_csv(VIEWER, "prj_obc"))))
        self.assertEqual([(r["board"], r["rename_to"], r["connectors"]) for r in rows], [("OBC-A", "SPI_CLK", "J7.3")])
        self.assertEqual(self.service.systems_using_project(VIEWER, "prj_pwr")["systems"][0]["renames"], [])

    def test_the_report_lists_proposals(self) -> None:
        from openpyxl import load_workbook

        self.propose()
        content, _name, _version = self.service.report(DESIGNER, self.sid, "xlsx")
        sheet = load_workbook(io.BytesIO(content))["Renames"]
        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[1][1:5], [self.spi_row()["net_a"][0], "SPI_CLK", "1", "open"])


class RenameApiTest(RenameCase):
    def setUp(self) -> None:
        super().setUp()
        patcher = mock.patch.object(service_module, "service", self.service)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.app = FastAPI()
        self.app.include_router(systems_api.router, prefix="/api/systems")

    def test_propose_withdraw_and_the_board_page_over_http(self) -> None:
        body = {"instanceId": self.instances["OBC-A"], "net": self.spi_row()["net_a"][0], "name": "SPI_CLK"}
        path = f"/api/systems/{self.sid}/renames"
        self.assertEqual(_request(self.app, "POST", path, body=body).status, 428)
        created = _request(self.app, "POST", path, body=body, headers={"If-Match": etag(self.sid, self.version())})
        self.assertEqual(created.status, 201)
        listed = _request(self.app, "GET", "/api/systems/by-project/prj_obc", user="viewer")
        self.assertEqual(listed.json["systems"][0]["renames"][0]["id"], created.json["id"])
        exported = _request(self.app, "GET", "/api/systems/by-project/prj_obc/renames.csv", user="viewer")
        self.assertEqual(exported.status, 200)
        self.assertIn("SPI_CLK", exported.text)
        gone = _request(self.app, "DELETE", f"{path}/{created.json['id']}", headers={"If-Match": created.headers["etag"]})
        self.assertEqual(gone.status, 204)


if __name__ == "__main__":
    unittest.main()
